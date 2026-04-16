import pandas as pd
import openpyxl

try:
    wb = openpyxl.load_workbook('c:\\xampp\\htdocs\\candidate-test\\excel\\beam-analysis.xlsx', data_only=False)
    for sheet in wb.sheetnames:
        print(f"Sheet: {sheet}")
        ws = wb[sheet]
        for row in range(1, 20):
            row_vals = []
            for col in range(1, 10):
                cell = ws.cell(row=row, column=col)
                row_vals.append(str(cell.value) if cell.value else '')
            if any(row_vals):
                print(f"Row {row}: {row_vals}")
except Exception as e:
    print(e)
